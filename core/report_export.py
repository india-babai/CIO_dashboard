"""
=============================================================================
 REPORT EXPORT  —  "Download full report": everything in one .xlsx
=============================================================================

WHAT THIS FILE DOES
    Writes ONE Excel workbook containing every table and every chart's
    underlying data from the whole app, so a user can take whatever they need
    without going back to the screen.

    SHEETS IT PRODUCES  (each is skipped silently if its data was not supplied)

      README                what the file contains, and the settings used
      Allocation All RPs    the house table exactly as page 1 shows it
      Allocation RP1..RP5   one sheet per profile, stacked house format
      Risk-Return Metrics   every profile x scenario
      CMA                   expected return + volatility actually used
      Correlation Matrix    the full matrix
      Frontier Curve        the efficient frontier points  (+ Excel chart)
      Frontier vs Model     each profile's risk/return and its return gap
      Backtest Growth       month-end index levels         (+ Excel chart)
      Backtest Drawdown     month-end drawdown             (+ Excel chart)
      Backtest Rolling Ret  rolling window return
      Backtest Rolling Vol  rolling window volatility
      Backtest Metrics      the full backtest metrics table
      Price History         the daily prices the backtest ran on

    Native Excel line charts are embedded on the Growth, Drawdown and Frontier
    sheets, so the workbook has real charts rather than only numbers.

WHY THE BACKTEST SHEETS ARE MONTH-END
    The backtest runs on ~5,000 daily points. Writing those to Excel makes a
    slow, heavy file and an unreadable chart. The series are resampled to
    month-end for these sheets; the full daily prices are on "Price History"
    if you need them.

WHERE TO CHANGE WHAT
    * Add a sheet          -> write a _write_*() helper and call it from
                              build_workbook() below
    * Sheet contents       -> the matching _write_*() helper
    * Chart type / anchor  -> _add_line_chart() / _add_scatter_chart()
    * Which model is used  -> ui/report_builder.py (it gathers the inputs)

NO STREAMLIT IN THIS FILE - it is handed plain DataFrames and returns bytes.
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.utils import get_column_letter

from core.config import Config
from core.report_tables import all_columns, stacked_block

MAX_SHEET_NAME = 31


def _safe_sheet_name(name: str) -> str:
    """Excel sheet names: 31 chars, and none of  : \\ / ? * [ ]"""
    for bad in ':\\/?*[]':
        name = name.replace(bad, '-')
    return name[:MAX_SHEET_NAME]


def _autosize(worksheet, frame: pd.DataFrame, index_width: int = 22) -> None:
    """Rough column widths so the sheet is readable without fiddling."""
    worksheet.column_dimensions["A"].width = index_width
    for position, column in enumerate(frame.columns, start=2):
        header = str(column[-1] if isinstance(column, tuple) else column)
        worksheet.column_dimensions[get_column_letter(position)].width = \
            max(11, min(28, len(header) + 3))


def _write(writer, frame: pd.DataFrame, sheet: str, index: bool = True,
           index_width: int = 22):
    """Write a frame and return its worksheet (or None if the frame is empty)."""
    if frame is None or frame.empty:
        return None
    sheet = _safe_sheet_name(sheet)
    frame.to_excel(writer, sheet_name=sheet, index=index)
    worksheet = writer.sheets[sheet]
    _autosize(worksheet, frame, index_width)
    return worksheet


def _add_line_chart(worksheet, frame: pd.DataFrame, title: str,
                    y_title: str, anchor: str) -> None:
    """Embed a native Excel line chart covering every column of `frame`."""
    rows, cols = frame.shape
    if rows < 2 or cols < 1:
        return
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.height, chart.width = 9, 22
    chart.style = 2
    data = Reference(worksheet, min_col=2, max_col=1 + cols,
                     min_row=1, max_row=1 + rows)
    categories = Reference(worksheet, min_col=1, min_row=2, max_row=1 + rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    for series in chart.series:
        series.smooth = False
    worksheet.add_chart(chart, anchor)


def _add_scatter_chart(worksheet, rows: int, title: str,
                       x_title: str, y_title: str, anchor: str) -> None:
    """Scatter chart of column A (x) against column B (y)."""
    if rows < 2:
        return
    chart = ScatterChart()
    chart.title = title
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.height, chart.width = 9, 18
    chart.style = 2
    x = Reference(worksheet, min_col=1, min_row=2, max_row=1 + rows)
    y = Reference(worksheet, min_col=2, min_row=1, max_row=1 + rows)
    series = Series(y, x, title_from_data=True)
    series.marker.symbol = "none"
    chart.series.append(series)
    worksheet.add_chart(chart, anchor)


def _month_end(series: pd.Series) -> pd.Series:
    """Resample a daily series to month-end so Excel stays light."""
    if series is None or series.empty:
        return pd.Series(dtype=float)
    return series.resample("ME").last()


def _frame_from_series(series_by_name: dict) -> pd.DataFrame:
    """{name: series} -> one DataFrame, month-end, columns in the given order."""
    usable = {name: _month_end(series)
              for name, series in (series_by_name or {}).items()
              if series is not None and not series.empty}
    if not usable:
        return pd.DataFrame()
    frame = pd.DataFrame(usable)
    frame.index.name = "Month end"
    return frame.round(4)


# --------------------------------------------------------------------------- #
# the cover sheet                                                            #
# --------------------------------------------------------------------------- #
def _readme_frame(cfg: Config, model_name: str, currency: str,
                  included: list[str], notes: list[str]) -> pd.DataFrame:
    newer, older = cfg.scenario_change
    rows = [
        ("CIO Model Portfolio Dashboard — full report", ""),
        ("", ""),
        ("Model portfolio", model_name),
        ("Currency / CMA used", currency),
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("SETTINGS USED", ""),
        ("Scenarios", " / ".join(cfg.scenarios)),
        ("Change column", f"{newer} minus {older}"),
        ("Tracking error benchmark", cfg.scenario_benchmark),
        ("Backtest rebalancing", str(cfg.backtest.get("rebalance", "M"))),
        ("Backtest risk-free rate %", str(cfg.backtest.get("risk_free", 0.0))),
        ("VaR / CVaR confidence", str(cfg.backtest.get("var_confidence", 0.95))),
        ("Frontier points solved", str(cfg.frontier.get("points", 120))),
        ("", ""),
        ("SHEETS IN THIS FILE", ""),
    ]
    rows += [(f"  {name}", "") for name in included]
    if notes:
        rows += [("", ""), ("NOTES", "")]
        rows += [(f"  {note}", "") for note in notes]

    frame = pd.DataFrame(rows, columns=["", " "])
    return frame.set_index("")


# --------------------------------------------------------------------------- #
# the builder                                                                #
# --------------------------------------------------------------------------- #
def build_workbook(cfg: Config, *, model_name: str, currency: str,
                   weights_by_rp: dict, cma: pd.DataFrame, corr: pd.DataFrame,
                   house_table: pd.DataFrame | None = None,
                   metrics_all: pd.DataFrame | None = None,
                   frontier_points: pd.DataFrame | None = None,
                   frontier_gaps: pd.DataFrame | None = None,
                   backtest_growth: dict | None = None,
                   backtest_drawdown: dict | None = None,
                   backtest_rolling_return: dict | None = None,
                   backtest_rolling_vol: dict | None = None,
                   backtest_metrics: pd.DataFrame | None = None,
                   prices: pd.DataFrame | None = None,
                   notes: list[str] | None = None) -> bytes:
    """
    Assemble the workbook. Every argument after `corr` is optional - whatever
    is not supplied simply does not get a sheet, so the caller can build a
    partial report (e.g. when there is no price history).
    """
    buffer = io.BytesIO()
    included: list[str] = []
    notes = list(notes or [])

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # the cover is written last (we need the sheet list first) but must be
        # the FIRST tab, so reserve it now with a placeholder and rewrite later
        placeholder = pd.DataFrame({"": ["building..."]}).set_index("")
        placeholder.to_excel(writer, sheet_name="README")

        # ---- allocations ---------------------------------------------- #
        if house_table is not None and not house_table.empty:
            _write(writer, house_table, "Allocation All RPs", index_width=26)
            included.append("Allocation All RPs — the house table from page 1")

        for rp, weights in (weights_by_rp or {}).items():
            sheet = f"Allocation {cfg.rp_short(rp)}"
            _write(writer, stacked_block(cfg, weights, cma, corr), sheet,
                   index_width=26)
        if weights_by_rp:
            included.append("Allocation RP1..RP5 — one sheet per risk profile")

        if metrics_all is not None and not metrics_all.empty:
            _write(writer, metrics_all, "Risk-Return Metrics", index_width=24)
            included.append("Risk-Return Metrics — every profile x scenario")

        # ---- the inputs ------------------------------------------------ #
        if cma is not None and not cma.empty:
            cma_sheet = cma.set_index("code") if "code" in cma.columns else cma
            _write(writer, cma_sheet.round(4), "CMA", index_width=14)
            included.append("CMA — expected return and volatility used")

        if corr is not None and not corr.empty:
            _write(writer, corr.round(4), "Correlation Matrix", index_width=14)
            included.append("Correlation Matrix")

        # ---- frontier --------------------------------------------------- #
        if frontier_points is not None and not frontier_points.empty:
            sheet = _write(writer, frontier_points.round(4), "Frontier Curve",
                           index=False, index_width=14)
            if sheet is not None:
                _add_scatter_chart(sheet, len(frontier_points),
                                   "Efficient frontier",
                                   "Volatility %", "Expected return %", "E2")
            included.append("Frontier Curve — solved points (with chart)")

        if frontier_gaps is not None and not frontier_gaps.empty:
            _write(writer, frontier_gaps.round(4), "Frontier vs Model",
                   index=False, index_width=18)
            included.append("Frontier vs Model — return gap per profile")

        # ---- backtest ---------------------------------------------------- #
        growth = _frame_from_series(backtest_growth)
        if not growth.empty:
            sheet = _write(writer, growth, "Backtest Growth", index_width=14)
            if sheet is not None:
                _add_line_chart(sheet, growth, "Growth of 100 (month end)",
                                "Index", f"{get_column_letter(len(growth.columns)+3)}2")
            included.append("Backtest Growth — month-end index (with chart)")

        drawdown = _frame_from_series(backtest_drawdown)
        if not drawdown.empty:
            sheet = _write(writer, drawdown, "Backtest Drawdown", index_width=14)
            if sheet is not None:
                _add_line_chart(sheet, drawdown, "Drawdown (month end)",
                                "%", f"{get_column_letter(len(drawdown.columns)+3)}2")
            included.append("Backtest Drawdown — month-end (with chart)")

        rolling_return = _frame_from_series(backtest_rolling_return)
        if not rolling_return.empty:
            _write(writer, rolling_return, "Backtest Rolling Ret", index_width=14)
            included.append("Backtest Rolling Ret — rolling window return")

        rolling_vol = _frame_from_series(backtest_rolling_vol)
        if not rolling_vol.empty:
            _write(writer, rolling_vol, "Backtest Rolling Vol", index_width=14)
            included.append("Backtest Rolling Vol — rolling window volatility")

        if backtest_metrics is not None and not backtest_metrics.empty:
            _write(writer, backtest_metrics.round(4), "Backtest Metrics",
                   index_width=24)
            included.append("Backtest Metrics — the full metrics table")

        if prices is not None and not prices.empty:
            _write(writer, prices.round(4), "Price History", index_width=14)
            included.append(f"Price History — {len(prices):,} daily rows")
            notes.append("Price History holds the full daily series; the "
                         "backtest sheets are resampled to month end.")

        # ---- rewrite the cover now that we know what went in ------------- #
        del writer.book["README"]
        readme = _readme_frame(cfg, model_name, currency, included, notes)
        readme.to_excel(writer, sheet_name="README")
        cover = writer.sheets["README"]
        cover.column_dimensions["A"].width = 34
        cover.column_dimensions["B"].width = 62
        writer.book.move_sheet("README", offset=-len(writer.book.sheetnames) + 1)

    return buffer.getvalue()
